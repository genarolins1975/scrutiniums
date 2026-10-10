from pipeline.eficiencia_alimentar.textos import normalizar_textos
"""MUNIC2024 SAN: municipal self-reported institutional presence and acquisition actions. Not SISAN accession."""
import json,pathlib,gzip,hashlib,datetime,tempfile,urllib.request,sys,collections
import openpyxl
ROOT=pathlib.Path(__file__).resolve().parents[2];OUT=ROOT/'public/eficiencia/seguranca-alimentar';URL='https://ftp.ibge.gov.br/Perfil_Municipios/Seguranca_Alimentar_2024/Base_de_dados_Seguranca_Alimentar_MUNIC.xlsx'
# id, label, worksheet, column, quantity(parent existence), reference, section
CONFIG=[
 ('munic-gestor','Municípios com órgão gestor de SAN','Órgão gestor perfil do titular','SORG01',None,'2024','governanca'),
 ('munic-lei','Municípios com lei de SAN','Legislação inst gestão particip','SLIG01',None,'2024','governanca'),
 ('munic-plano','Municípios com plano de SAN','Legislação inst gestão particip','SLIG02',None,'2024','governanca'),
 ('munic-conselho','Municípios com conselho de SAN','Legislação inst gestão particip','SLIG05',None,'2024','governanca'),
 ('munic-camara','Municípios com câmara intersetorial de SAN','Legislação inst gestão particip','SLIG13',None,'2024','governanca'),
 ('munic-fundo','Municípios com fundo de SAN','Legislação inst gestão particip','SLIG21',None,'2024','governanca'),
 ('munic-restaurante','Municípios com restaurante popular sob gestão municipal','Equipamentos','SEQP14',None,'2024','equipamentos'),
 ('munic-restaurante-quantidade','Restaurantes populares sob gestão municipal','Equipamentos','SEQP141','SEQP14','2024','equipamentos'),
 ('munic-cozinha','Municípios com cozinha comunitária sob gestão municipal','Equipamentos','SEQP29',None,'2024','equipamentos'),
 ('munic-cozinha-quantidade','Cozinhas comunitárias sob gestão municipal','Equipamentos','SEQP291','SEQP29','2024','equipamentos'),
 ('munic-banco','Municípios com banco de alimentos sob gestão municipal','Equipamentos','SEQP44',None,'2024','equipamentos'),
 ('munic-banco-quantidade','Bancos de alimentos sob gestão municipal','Equipamentos','SEQP441','SEQP44','2024','equipamentos'),
 ('munic-central','Municípios com central de recebimento da agricultura familiar','Equipamentos','SEQP57',None,'2024','equipamentos'),
 ('munic-central-quantidade','Centrais de recebimento da agricultura familiar','Equipamentos','SEQP571','SEQP57','2024','equipamentos'),
 ('munic-paa','Municípios que declararam aquisição via PAA','Ações','SACO07',None,'2023','aquisicao'),
 ('munic-pnae','Municípios que declararam aquisição via PNAE','Ações','SACO10',None,'2023','aquisicao'),
 ('munic-acesso-alimentos','Municípios com ações de acesso a alimentos','Ações','SACO02',None,'2023','aquisicao'),
 ('munic-cestas','Municípios com distribuição de cestas básicas','Ações','SACO0211',None,'2023','aquisicao'),
 ('munic-refeicoes-prontas','Municípios com oferta de refeições prontas','Ações','SACO0212',None,'2023','aquisicao'),
 ('munic-ticket','Municípios com ticket ou vale-alimentação','Ações','SACO0213',None,'2023','aquisicao'),
 ('munic-beneficio-monetario','Municípios com oferta de benefício monetário','Ações','SACO0214',None,'2023','aquisicao')]
def binary(value,col):
 if col=='SORG01':return None if value in [None,'Recusa','Não informou'] else 0 if value=='Não possui estrutura' else 1
 if col=='SLIG01':return 1 if value=='Possui lei' else 0 if value in ['Não possui','A lei está em trâmite'] else None
 return 1 if value=='Sim' else 0 if value=='Não' else None

def build():
 p=OUT/'brutos/munic-san-2024.xlsx.gz'
 if '--coletar' in sys.argv or not p.exists():p.parent.mkdir(parents=True,exist_ok=True);p.write_bytes(gzip.compress(urllib.request.urlopen(URL,timeout=60).read(),mtime=0))
 raw=gzip.decompress(p.read_bytes())
 with tempfile.TemporaryDirectory() as folder:
  file=pathlib.Path(folder)/'munic.xlsx';file.write_bytes(raw);w=openpyxl.load_workbook(file,read_only=True,data_only=True);sheets={sh.title:[dict(zip(next(sh.iter_rows(values_only=True)),r)) for r in list(sh.iter_rows(values_only=True))[1:] if isinstance(r[0],int)] for sh in w if sh.title!='Dicionário'}
  if any(len(v)!=5570 for v in sheets.values()):raise ValueError('Unexpected MUNIC territorial coverage')
  dic={str(row[6]):str(row[5]) for row in w['Dicionário'].iter_rows(values_only=True) if len(row)>6 and row[6] is not None}
 ebia=json.loads(OUT.joinpath('snapshot.json').read_text());ts={t['id']:t for t in ebia['territorios']}
 for r in sheets['Órgão gestor perfil do titular']:ts[str(r['CodMun'])]={'id':str(r['CodMun']),'nome':r['Mun'],'nivel':'municipio','uf':r['Sigla UF']}
 source={'id':'munic-san-2024','nome':'IBGE · MUNIC2024 · Suplemento Segurança Alimentar e Nutricional','url':URL,'consulta':URL,'periodo':{'frequencia':'retrato','inicio':2023,'fim':2024},'universo':'Declarações das prefeituras dos5570municípios na MUNIC2024. Instituições e equipamentos na pesquisa; ações de aquisição e promoção do acesso referem-se a2023.','limitacoes':['Pesquisa de informações prestadas pela prefeitura, não censo de beneficiários nem aferição direta da oferta em funcionamento.','Existência de conselho, plano, órgão gestor ou equipamento não significa adesão ao SISAN, cumprimento das obrigações ou qualidade do serviço.','Quantidades de equipamentos de responsabilidade municipal incluem parcerias/convênios conforme definição do questionário; não cobrem todo equipamento estadual ou privado no território.','Recusa, não informou e não sabe informar são ausências, nunca zeros. Traço condicional só vira zero quando a pergunta anterior declara explicitamente inexistência.','Contagens Brasil, regiões eUF são somas municipais de respostas conhecidas; cobertura de respondentes acompanha cada agregado. Não somar esses níveis entre si.','PAA/PNAE e ações de acesso em2023: presença declarada, não valor financeiro executado, volume comprado, atendimento ou resultado causal.'],'capturadoEm':datetime.datetime.now(datetime.timezone.utc).isoformat(),'sha256':hashlib.sha256(raw).hexdigest(),'bruto':'/eficiencia/seguranca-alimentar/brutos/munic-san-2024.xlsx.gz','transformacoes':['Presença:Sim=1 eNão=0; recusas/desconhecidos/null preservados como ausência.','Quantidade:número original; pergunta de quantidade não aplicável por inexistência explícita recebe zero calculado, identificado.','Somas das respostas municipais conhecidas; totais de municípios do território e municípios com dado são preservados.']}
 datasets={k:{'versao':1,'capturadoEm':source['capturadoEm'],'fontes':[source],'territorios':list(ts.values()),'indicadores':[],'observacoes':[],'notas':source['limitacoes']} for k in ['sisan','aquisicao']}
 regionmap={'1':'regiao-1','2':'regiao-2','3':'regiao-3','4':'regiao-4','5':'regiao-5'}
 ufreg={'11':'1','12':'1','13':'1','14':'1','15':'1','16':'1','17':'1','21':'2','22':'2','23':'2','24':'2','25':'2','26':'2','27':'2','28':'2','29':'2','31':'3','32':'3','33':'3','35':'3','41':'4','42':'4','43':'4','50':'5','51':'5','52':'5','53':'5'}
 for id,name,title,col,parent,period,section in CONFIG:
  target=datasets['aquisicao' if section=='aquisicao' else 'sisan'];unit='equipamentos' if parent else 'municípios';target['indicadores'].append({'id':id,'nome':name,'unidade':unit,'fonteId':source['id'],'universo':f'MUNIC2024: declaração municipal, campo {col}, {title}.','limitacao':'Resposta autodeclarada. No município,presença=1/ausência explicitamente declarada=0; Brasil/UF/região mostram soma com cobertura, não taxa.' if not parent else 'Quantidade declarada; ausência explícita do equipamento implica zero. Falta de resposta não é zero.','frequencia':'retrato','categoria':section,'perimetro':'declaração municipal','campoFonte':col})
  aggregates={}
  for r in sheets[title]:
   tid=str(r['CodMun']);original=r.get(col);derived=False;note='Resposta municipal: '+str(original)
   if parent:
    val=float(original) if isinstance(original,(int,float)) else None
    if val is None and binary(r.get(parent),parent)==0:val=0;derived=True;note='Zero decorrente de inexistência explicitamente declarada; quantidade não aplicável.'
   else:
    val=binary(original,col)
    if col.startswith('SACO021') and original=='-' and r.get('SACO02')=='Não':val=0;derived=True;note='Zero decorrente de inexistência explícita de ação de promoção do acesso.'
   base={'indicadorId':id,'territorioId':tid,'periodo':period,'periodoNome':period+' · MUNIC2024','grupo':'Total','dimensao':'total','valor':val,'status':'ausente' if val is None else 'calculado' if derived else 'observado','cv':None,'fonteId':source['id'],'nota':note,'valorOriginal':original,'campoFonte':col}
   target['observacoes'].append(base)
   for aid in ['1',tid[:2],regionmap[ufreg[tid[:2]]]]:
    if aid not in aggregates:aggregates[aid]={'sum':0,'known':0,'all':0}
    aggregates[aid]['all']+=1
    if val is not None:aggregates[aid]['sum']+=val;aggregates[aid]['known']+=1
  for aid,a in aggregates.items():target['observacoes'].append({'indicadorId':id,'territorioId':aid,'periodo':period,'periodoNome':period+' · MUNIC2024','grupo':'Total','dimensao':'total','valor':a['sum'] if a['known'] else None,'status':'calculado' if a['known'] else 'ausente','cv':None,'fonteId':source['id'],'nota':f'Soma das respostas municipais conhecidas: {a["known"]} de {a["all"]} municípios. Ausências não imputadas.','municipiosComDado':a['known'],'municipiosUniverso':a['all'],'campoFonte':col})
 for key,d in datasets.items():OUT.joinpath(key+'.json').write_text(json.dumps(normalizar_textos(d),ensure_ascii=False,separators=(',',':')));print(key,len(d['indicadores']),'indicadores',len(d['observacoes']),'observações')
if __name__=='__main__':build()
