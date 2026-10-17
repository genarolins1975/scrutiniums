"""Comparação internacional (ferramenta transversal): estimativas de água e esgotamento do JMP (OMS e UNICEF).

Fonte: WHO/UNICEF Joint Monitoring Programme for Water Supply, Sanitation and Hygiene (washdata.org), base de
domicílios, arquivo mundial com as estimativas de todos os países, áreas e territórios, 2000 a 2024. A pasta de
trabalho tem abas formatadas e abas de dados sem arredondamento (`wat` e `san`); só as segundas são lidas. O
arquivo (25 MB) não é versionado: o seed guarda as colunas usadas, com o sha256 e o tamanho do original.

Nomes de países em português vêm da API de países do IBGE (ISO 3166-1 alfa-3). Áreas que o IBGE não lista
(territórios) ficam com o nome em inglês do JMP, marcado como tal. Nenhuma tradução é feita à mão.

O JMP é uma modelagem sobre censos, pesquisas domiciliares e registros administrativos de cada país, com
definições próprias (fonte melhorada, tempo de coleta, "gerenciado com segurança"); não é medida do SINISA nem
do Censo. A gold guarda os valores como o JMP os publica, em percentual da população; posição e comparações
entram na camada de apresentação.

Coleta: python3 -m pipeline.observatorio_saneamento.jmp --coletar
"""
import argparse
import csv
import datetime as dt
import gzip
import hashlib
import io
import json
import re
import urllib.request
from pathlib import Path

AQUI = Path(__file__).resolve().parent
SEED = AQUI / "seed" / "jmp"
URL_JMP = "https://washdata.org/data/country/WLD/household/download"
URL_BRASIL = "https://washdata.org/data/country/BRA/household/download"
URL_PAISES_IBGE = "https://servicodados.ibge.gov.br/api/v1/paises/all"
UA = "ObservatorioBrasileiroDeCredito coleta jmp"
TOLERANCIA_ESCADA = 0.05  # soma dos degraus do serviço, em pontos percentuais (valores sem arredondamento)

IDENT = ["name", "iso3", "year", "pop_t", "prop_u", "region_sdg", "region_unicef_programme", "region_who", "region_income"]
AREAS = ("t", "u", "r")  # total, urbana, rural
ESCADA_AGUA = ["basal", "lim", "unimp", "ns"]  # pelo menos básico, limitado, não melhorado, água superficial
ESCADA_ESGOTO = ["basal", "lim", "unimp", "ns"]  # pelo menos básico, limitado, não melhorado, defecação a céu aberto
COLUNAS_AGUA = [f"wat_{x}_{a}" for a in AREAS for x in ESCADA_AGUA + ["sm"]] + ["wat_pip_t", "wat_imp_prem_t", "wat_imp_av_t", "wat_imp_qual_t"]
COLUNAS_ESGOTO = [f"san_{x}_{a}" for a in AREAS for x in ESCADA_ESGOTO + ["sm"]] + [
    "san_sew_t", "san_sep_t", "san_sew_sm_t", "san_sdo_sm_t", "san_fst_sm_t", "san_othimp_t"]
ABAS = {"agua": ("wat", COLUNAS_AGUA), "esgoto": ("san", COLUNAS_ESGOTO)}


def sha(b):
    return hashlib.sha256(b).hexdigest()


def numero(v):
    """Célula da aba de dados → float ou None. Vazio, '-' e texto não numérico são ausência."""
    if v is None or isinstance(v, bool):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------------------
# coleta
# ---------------------------------------------------------------------------

def _baixar(url, timeout=600):
    with urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": UA, "Accept-Encoding": "identity"}), timeout=timeout) as r:
        return r.read(), r.headers.get("Last-Modified"), r.url


def _gravar_csv(nome, cabecalho, linhas):
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\n")
    w.writerow(cabecalho)
    for l in linhas:
        w.writerow(["" if v is None else v for v in l])
    bruto = buf.getvalue().encode("utf-8")
    (SEED / f"{nome}.gz").write_bytes(gzip.compress(bruto, mtime=0))
    return {"arquivo": f"{nome}.gz", "sha256": sha(bruto), "bytes": len(bruto), "linhas": len(linhas)}


MESES = {m: i for i, m in enumerate(
    ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"], 1)}


def mes_da_edicao(edicao):
    """"August 2025" → "2025-08"."""
    nome, ano = edicao.split()
    return f"{ano}-{MESES[nome]:02d}"


def edicao_do_arquivo(textos):
    """"Updated August 2025" e "(2000-2024)" da aba Introduction."""
    junto = " ".join(textos)
    atualizado = re.search(r"Updated\s+([A-Za-z]+\s+\d{4})", junto)
    periodo = re.search(r"\((\d{4})\s*-\s*(\d{4})\)", junto)
    if not atualizado or not periodo:
        raise ValueError("aba Introduction sem a data de atualização ou o período: arquivo do JMP em formato novo")
    return atualizado.group(1), (int(periodo.group(1)), int(periodo.group(2)))


FONTE_BRASIL = re.compile(r"^\[?(BRA_(\d{4})_([A-Za-z0-9]+))\]?$")


def fontes_do_brasil(linhas):
    """Fontes que o JMP lista para o Brasil na aba Data Summary: código (BRA_2022_CEN), sigla, tipo e ano.

    A aba também marca entre colchetes valores que o JMP não usou; isso é por valor, não por fonte, e não é lido aqui.
    Linha fora do padrão de código é ignorada."""
    saida = []
    for primeira, tipo in linhas:
        m = FONTE_BRASIL.match(str(primeira or "").strip())
        if m:
            saida.append({"codigo": m.group(1), "ano": int(m.group(2)), "sigla": m.group(3), "tipo": str(tipo or "")})
    return saida


def coletar():
    import openpyxl  # dependência só do coletor

    SEED.mkdir(parents=True, exist_ok=True)
    xlsx, modificado, url_final = _baixar(URL_JMP)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx), read_only=True, data_only=True)
    intro = [str(c) for row in wb["Introduction"].iter_rows(values_only=True) for c in row if c]
    atualizado, periodo = edicao_do_arquivo(intro)
    arquivos = {}
    for chave, (aba, colunas) in ABAS.items():
        linhas = wb[aba].iter_rows(values_only=True)
        cab = list(next(linhas))
        faltam = [c for c in IDENT + colunas if c not in cab]
        if faltam:
            raise ValueError(f"aba {aba} do JMP sem as colunas {faltam}")
        idx = [cab.index(c) for c in IDENT + colunas]
        saida = [[r[i] for i in idx] for r in linhas if r[0] is not None]
        arquivos[chave] = _gravar_csv(f"{chave}.csv", IDENT + colunas, saida)
    brasil_xlsx, brasil_modificado, _ = _baixar(URL_BRASIL)
    wb_br = openpyxl.load_workbook(io.BytesIO(brasil_xlsx), read_only=True, data_only=True)
    fontes = fontes_do_brasil((r[0], r[1]) for r in wb_br["Data Summary"].iter_rows(values_only=True))
    if not fontes:
        raise ValueError("aba Data Summary do arquivo do Brasil sem nenhuma fonte reconhecida")
    (SEED / "brasil_fontes.json.gz").write_bytes(gzip.compress(json.dumps(fontes, ensure_ascii=False).encode("utf-8"), mtime=0))
    paises_bruto, _, _ = _baixar(URL_PAISES_IBGE, 120)
    paises = json.loads(gzip.decompress(paises_bruto) if paises_bruto[:2] == b"\x1f\x8b" else paises_bruto)
    reduzido = [{"iso3": p["id"]["ISO-3166-1-ALPHA-3"], "nome": p["nome"]["abreviado"]} for p in paises]
    (SEED / "paises_ibge.json.gz").write_bytes(gzip.compress(json.dumps(reduzido, ensure_ascii=False, sort_keys=True).encode("utf-8"), mtime=0))
    manifesto = {
        "capturado_em": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
        "url": URL_JMP,
        "url_final": url_final,
        "edicao": atualizado,
        "periodo": list(periodo),
        "xlsx": {"sha256": sha(xlsx), "bytes": len(xlsx), "last_modified": modificado},
        "brasil_xlsx": {"url": URL_BRASIL, "sha256": sha(brasil_xlsx), "bytes": len(brasil_xlsx), "last_modified": brasil_modificado, "fontes": len(fontes)},
        "paises_ibge": {"url": URL_PAISES_IBGE, "sha256": sha(paises_bruto), "bytes": len(paises_bruto), "paises": len(reduzido)},
        "arquivos": arquivos,
    }
    (SEED / "manifesto.json").write_text(json.dumps(manifesto, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"JMP {atualizado}, {periodo[0]} a {periodo[1]}:", {k: v["linhas"] for k, v in arquivos.items()}, flush=True)


# ---------------------------------------------------------------------------
# construção da gold
# ---------------------------------------------------------------------------

def _ler_csv(manifesto, chave):
    a = manifesto["arquivos"][chave]
    bruto = gzip.decompress((SEED / a["arquivo"]).read_bytes())
    if sha(bruto) != a["sha256"]:
        raise ValueError(f"seed do JMP alterado: {a['arquivo']}")
    return list(csv.DictReader(io.StringIO(bruto.decode("utf-8")), delimiter=";"))


def _bloco(l, prefixo, extras, area):
    """Valores de uma área (t, u, r) como o JMP publica. Degrau ausente é ausência; não vira zero."""
    r = {
        "basica": numero(l[f"{prefixo}_basal_{area}"]),
        "limitada": numero(l[f"{prefixo}_lim_{area}"]),
        "nao_melhorada": numero(l[f"{prefixo}_unimp_{area}"]),
        "ruim": numero(l[f"{prefixo}_ns_{area}"]),  # água superficial (água) ou defecação a céu aberto (esgoto)
        "gerenciada": numero(l[f"{prefixo}_sm_{area}"]),
    }
    if area == "t":
        r.update({k: numero(l[f"{prefixo}_{c}_t"]) for k, c in extras})
    return r


EXTRAS_AGUA = [("encanada", "pip"), ("no_domicilio", "imp_prem"), ("disponivel", "imp_av"), ("sem_contaminacao", "imp_qual")]
EXTRAS_ESGOTO = [("rede", "sew"), ("fossa_septica", "sep"), ("rede_tratada", "sew_sm"), ("fossa_esvaziada_tratada", "sdo_sm"),
                 ("fossa_sem_esvaziamento", "fst_sm"), ("outras_melhoradas", "othimp")]


def _conferir(linhas, prefixo, aba):
    """Degraus do serviço somam 100% e 'gerenciado' não passa de 'pelo menos básico', em toda linha com estimativa."""
    validas, maior = 0, 0.0
    for l in linhas:
        for a in AREAS:
            basica = numero(l[f"{prefixo}_basal_{a}"])
            if basica is None:
                continue
            soma = sum(numero(l[f"{prefixo}_{x}_{a}"]) or 0 for x in ("basal", "lim", "unimp", "ns"))
            validas += 1
            maior = max(maior, abs(soma - 100))
            if abs(soma - 100) > TOLERANCIA_ESCADA:
                raise ValueError(f"{aba}: {l['name']} {l['year']} ({a}): degraus somam {soma:.3f}%")
            sm = numero(l[f"{prefixo}_sm_{a}"])
            if sm is not None and sm > basica + 0.01:
                raise ValueError(f"{aba}: {l['name']} {l['year']} ({a}): gerenciado ({sm:.2f}) acima de pelo menos básico ({basica:.2f})")
    return {"estimativas": validas, "maior_desvio_da_soma": round(maior, 4)}


def construir():
    manifesto = json.loads((SEED / "manifesto.json").read_text(encoding="utf-8"))
    agua, esgoto = _ler_csv(manifesto, "agua"), _ler_csv(manifesto, "esgoto")
    ibge = {p["iso3"]: p["nome"] for p in json.loads(gzip.decompress((SEED / "paises_ibge.json.gz").read_bytes()))}
    conf = {"agua": _conferir(agua, "wat", "água"), "esgoto": _conferir(esgoto, "san", "esgoto")}
    fontes_brasil = json.loads(gzip.decompress((SEED / "brasil_fontes.json.gz").read_bytes()))
    ano = manifesto["periodo"][1]

    def por(linhas):
        d = {}
        for l in linhas:
            chave = (l["iso3"], int(l["year"]))
            if chave in d:
                raise ValueError(f"JMP: {chave} repetido")
            if not l["iso3"]:
                raise ValueError(f"JMP: {l['name']} sem código ISO alfa-3")
            d[chave] = l
        return d

    a, e = por(agua), por(esgoto)
    isos = sorted({k[0] for k in a} | {k[0] for k in e})

    def blocos(iso, y):
        la, le = a.get((iso, y)), e.get((iso, y))
        return {
            "populacao_mil": numero((la or le)["pop_t"]),
            "urbana_pct": numero((la or le)["prop_u"]),
            "ano": y,
            "agua": {ar: _bloco(la, "wat", EXTRAS_AGUA, ar) for ar in AREAS} if la else None,
            "esgoto": {ar: _bloco(le, "san", EXTRAS_ESGOTO, ar) for ar in AREAS} if le else None,
        }

    def linha_pais(iso):
        base = a.get((iso, ano)) or e.get((iso, ano))
        nome_jmp = base["name"].strip()
        return {
            "iso3": iso,
            "nome": ibge.get(iso, nome_jmp),
            "nome_origem": "IBGE" if iso in ibge else "JMP",
            "nome_jmp": nome_jmp,
            "regiao_ods": base["region_sdg"] or None,
            "regiao_unicef": base["region_unicef_programme"] or None,
            "regiao_oms": base["region_who"] or None,
            "renda": base["region_income"] or None,
        } | blocos(iso, ano)

    paises = [linha_pais(i) for i in isos if (i, ano) in a or (i, ano) in e]
    serie = [blocos("BRA", y) for y in range(manifesto["periodo"][0], ano + 1) if ("BRA", y) in a or ("BRA", y) in e]
    meta = {
        "fonte": manifesto["url"],
        "edicao": manifesto["edicao"],
        "edicao_mes": mes_da_edicao(manifesto["edicao"]),
        "periodo": manifesto["periodo"],
        "ano_comparado": ano,
        "capturado_em": manifesto["capturado_em"],
        "xlsx": manifesto["xlsx"],
        "paises_ibge": manifesto["paises_ibge"],
        "brasil_xlsx": manifesto["brasil_xlsx"],
        "brasil_fontes": fontes_brasil,
        "areas": len(paises),
        "com_agua_gerenciada": sum(1 for p in paises if p["agua"] and p["agua"]["t"]["gerenciada"] is not None),
        "com_esgoto_gerenciado": sum(1 for p in paises if p["esgoto"] and p["esgoto"]["t"]["gerenciada"] is not None),
        "sem_nome_ibge": [{"iso3": p["iso3"], "nome": p["nome_jmp"]} for p in paises if p["nome_origem"] == "JMP"],
        "conferencias": conf,
        "regras": {"tolerancia_escada_pct": TOLERANCIA_ESCADA},
    }
    return {"paises": paises, "brasil": serie}, meta


def linhas_csv(jmp):
    cab = ["iso3", "pais", "regiao_ods", "renda", "populacao_mil", "ano", "agua_basica_pct", "agua_limitada_pct", "agua_nao_melhorada_pct",
           "agua_superficial_pct", "agua_gerenciada_pct", "esgoto_basico_pct", "esgoto_limitado_pct", "esgoto_nao_melhorado_pct",
           "esgoto_ceu_aberto_pct", "esgoto_gerenciado_pct"]
    saida = []
    for p in jmp["paises"]:
        a, e = (p["agua"] or {}).get("t") or {}, (p["esgoto"] or {}).get("t") or {}
        saida.append([p["iso3"], p["nome"], p["regiao_ods"], p["renda"], p["populacao_mil"], p["ano"],
                      a.get("basica"), a.get("limitada"), a.get("nao_melhorada"), a.get("ruim"), a.get("gerenciada"),
                      e.get("basica"), e.get("limitada"), e.get("nao_melhorada"), e.get("ruim"), e.get("gerenciada")])
    return cab, saida


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--coletar", action="store_true", help="baixa o arquivo do JMP e a lista de países do IBGE e grava o seed (precisa de rede e de openpyxl)")
    if ap.parse_args().coletar:
        coletar()
    else:
        ap.print_help()
