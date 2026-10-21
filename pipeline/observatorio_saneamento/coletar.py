"""Coleta das fontes primárias e extração do recorte versionado (seed/).

Precisa de rede e de openpyxl (só para ler as planilhas do SINISA). A camada
gold (gold.py) não depende de nenhum dos dois: lê apenas seed/.

O recorte guarda, para cada planilha, o nome do arquivo dentro do ZIP, o
sha256 do ZIP e da planilha, a aba, a linha de cabeçalho e todas as colunas
codificadas (CAD, DFE, GTA, GTE, GFI, IAG, IES, IFA, IFE, IFB) linha a linha,
sem conversão: "Não calculado (...)" continua texto e ausência continua None.
Campos de observação livre (código terminado em 99) não entram.

Uso:
    python3 -m pipeline.observatorio_saneamento.coletar            # SINISA + SIDRA
    python3 -m pipeline.observatorio_saneamento.coletar --so-sidra
    python3 -m pipeline.observatorio_saneamento.coletar --zips DIR # ZIPs já baixados
"""
import argparse
import datetime as dt
import gzip
import hashlib
import io
import json
import os
import re
import sys
import time
import urllib.parse
import urllib.request
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

from pipeline.observatorio_saneamento import fontes  # noqa: E402

AQUI = os.path.dirname(os.path.abspath(__file__))
SEED = os.path.join(AQUI, "seed")
UA = "Scrutiniums/1.0 (observatorio do saneamento; dados abertos)"

UFS = {
    11: "RO", 12: "AC", 13: "AM", 14: "RR", 15: "PA", 16: "AP", 17: "TO",
    21: "MA", 22: "PI", 23: "CE", 24: "RN", 25: "PB", 26: "PE", 27: "AL", 28: "SE", 29: "BA",
    31: "MG", 32: "ES", 33: "RJ", 35: "SP", 41: "PR", 42: "SC", 43: "RS",
    50: "MS", 51: "MT", 52: "GO", 53: "DF",
}


def baixar(url, tentativas=4, timeout=300):
    erro = None
    for i in range(tentativas):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": UA})
            with urllib.request.urlopen(req, timeout=timeout) as r:
                b = r.read()
            # a API de metadados do IBGE responde gzip mesmo sem pedir
            return gzip.decompress(b) if b[:2] == b"\x1f\x8b" and not url.endswith((".zip", ".gz")) else b
        except Exception as e:  # noqa: BLE001 - qualquer falha de rede é retentada
            erro = e
            time.sleep(2 ** (i + 1))
    raise RuntimeError(f"falha ao baixar {url}: {erro}")


def sha256(b):
    return hashlib.sha256(b).hexdigest()


def gravar(nome, obj):
    os.makedirs(SEED, exist_ok=True)
    caminho = os.path.join(SEED, nome)
    dados = json.dumps(obj, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str).encode()
    # mtime fixo: o mesmo conteúdo produz o mesmo .gz byte a byte
    with open(caminho, "wb") as f:
        with gzip.GzipFile(fileobj=f, mode="wb", mtime=0, filename="") as g:
            g.write(dados)
    print(f"  {nome}: {len(dados) / 1e6:.1f} MB descomprimido")


def _decodificar_nome(n):
    # o ZIP do SINISA grava acentos como #U00e7; normaliza para comparar
    return re.sub(r"#U([0-9a-fA-F]{4})", lambda m: chr(int(m.group(1), 16)), n)


CODIGO = re.compile(r"^(CAD|DFE|GTA|GTE|GFI|IAG|IES|IFA|IFE|IFB|OGM)\d{4}[A-Z]?$", re.I)


def extrair_planilha(conteudo, nome_arquivo):
    import openpyxl  # dependência só do coletor

    wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    linhas = list(ws.iter_rows(values_only=True))
    cab = None
    for i, r in enumerate(linhas[:20]):
        if r and sum(1 for c in r if isinstance(c, str) and CODIGO.match(c.strip())) >= 3:
            cab = i
            break
    if cab is None:
        raise RuntimeError(f"cabeçalho de códigos não encontrado em {nome_arquivo}")
    # a planilha de ETEs grava os códigos em minúsculas (gte2101); normaliza
    codigos = [c.strip().upper() if isinstance(c, str) and CODIGO.match(c.strip()) else c.strip() if isinstance(c, str) else c
               for c in linhas[cab]]
    nomes = [re.sub(r"\s+", " ", c).strip() if isinstance(c, str) else c for c in linhas[cab - 2]]
    unidades = [re.sub(r"\s+", " ", c).strip() if isinstance(c, str) else c for c in linhas[cab - 1]]
    cabecalho_texto = [c for r in linhas[:cab - 2] for c in r if isinstance(c, str) and c.strip()]
    manter = [
        j for j, c in enumerate(codigos)
        if isinstance(c, str) and (c in ("cod_IBGE", "Macrorregião", "Município", "UF", "Capital", "RM/RIDE", "RM / RIDE",
                                         "Serviços", "Tipo de Serviço", "UF/Macrorregião", "Sigla") or (CODIGO.match(c) and not c.endswith("99")))
    ]
    colunas = [{"codigo": codigos[j], "nome": nomes[j], "unidade": unidades[j]} for j in manter]
    registros = []
    for k, r in enumerate(linhas[cab + 1:], start=cab + 2):
        if not r or all(c is None for c in r):
            continue
        vals = []
        for j in manter:
            v = r[j] if j < len(r) else None
            if isinstance(v, (dt.datetime, dt.date)):
                v = v.strftime("%Y-%m-%d")
            elif isinstance(v, str):
                v = v.strip()
            vals.append(v)
        registros.append({"linha": k, "v": vals})
    # nota metodológica (segunda aba), quando houver
    nota = []
    if len(wb.worksheets) > 1:
        for r in wb.worksheets[1].iter_rows(values_only=True):
            if r and any(c is not None for c in r):
                nota.append([re.sub(r"\s+", " ", c).strip() if isinstance(c, str) else c for c in r if c is not None])
    return {
        "aba": ws.title,
        "linha_codigos": cab + 1,
        "cabecalho": cabecalho_texto,
        "colunas": colunas,
        "registros": registros,
        "nota_metodologica": nota,
    }


def coletar_sinisa(dir_zips=None):
    capturado = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    zips = {}
    for modulo, url in fontes.SINISA_ZIPS.items():
        nome = url.rsplit("/", 1)[1]
        if dir_zips and os.path.exists(os.path.join(dir_zips, nome)):
            with open(os.path.join(dir_zips, nome), "rb") as f:
                b = f.read()
        else:
            print(f"baixando {url}")
            b = baixar(url)
        zips[modulo] = {"url": url, "arquivo": nome, "sha256": sha256(b), "bytes": len(b),
                        "zip": zipfile.ZipFile(io.BytesIO(b))}
    manifesto = {"capturado_em": capturado, "pagina": fontes.SINISA_PAGINA, "zips": {}, "planilhas": {}}
    for modulo, z in zips.items():
        manifesto["zips"][modulo] = {k: z[k] for k in ("url", "arquivo", "sha256", "bytes")}
    for chave, (modulo, padrao) in fontes.SINISA_PLANILHAS.items():
        z = zips[modulo]["zip"]
        alvo = [n for n in z.namelist() if padrao in _decodificar_nome(n) and n.endswith(".xlsx")]
        if len(alvo) != 1:
            raise RuntimeError(f"{chave}: esperado 1 arquivo para '{padrao}', achei {alvo}")
        conteudo = z.read(alvo[0])
        info = z.getinfo(alvo[0])
        print(f"extraindo {chave}")
        dados = extrair_planilha(conteudo, alvo[0])
        dados.update({
            "chave": chave,
            "modulo": modulo,
            "arquivo": _decodificar_nome(alvo[0]),
            "zip": zips[modulo]["arquivo"],
            "sha256": sha256(conteudo),
            "data_no_zip": "%04d-%02d-%02d" % info.date_time[:3],
        })
        manifesto["planilhas"][chave] = {k: dados[k] for k in ("arquivo", "zip", "sha256", "data_no_zip", "aba", "linha_codigos", "cabecalho")}
        manifesto["planilhas"][chave]["registros"] = len(dados["registros"])
        gravar(f"sinisa_{chave}.json.gz", dados)
    gravar("sinisa_manifesto.json.gz", manifesto)


def extrair_gestao_municipal(conteudo):
    """Planilha do módulo Gestão Municipal: aba GM (um município por linha) e as abas
    de detalhamento de prestadores e reguladores (uma linha por serviço e entidade)."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    limpa = lambda c: re.sub(r"\s+", " ", c).strip() if isinstance(c, str) else c  # noqa: E731
    saida = {}
    for ws in wb.worksheets:
        linhas = list(ws.iter_rows(values_only=True))
        cab = next(i for i, r in enumerate(linhas[:20]) if r and any(c in ("Cod_IBGE",) for c in r))
        cods = [limpa(c).replace("*", "").strip() if isinstance(c, str) else c for c in linhas[cab]]
        desc = [limpa(c) for c in linhas[cab + 1]]
        manter = [j for j, c in enumerate(cods) if c]
        registros = []
        for k, r in enumerate(linhas[cab + 2:], start=cab + 3):
            if not r or r[cods.index("Cod_IBGE")] in (None, ""):
                continue
            if not isinstance(r[cods.index("Cod_IBGE")], (int, float)) and not str(r[cods.index("Cod_IBGE")]).isdigit():
                continue
            vals = []
            for j in manter:
                v = r[j] if j < len(r) else None
                if isinstance(v, (dt.datetime, dt.date)):
                    v = v.strftime("%Y-%m-%d")
                vals.append(limpa(v))
            registros.append({"linha": k, "v": vals})
        saida[ws.title] = {
            "colunas": [{"codigo": cods[j], "nome": desc[j]} for j in manter],
            "registros": registros,
        }
    return saida


def coletar_gestao_municipal(arquivo_local=None):
    if arquivo_local and os.path.exists(arquivo_local):
        with open(arquivo_local, "rb") as f:
            b = f.read()
    else:
        b = baixar(fontes.SINISA_GESTAO_MUNICIPAL)
    dados = extrair_gestao_municipal(b)
    gravar("sinisa_gestao_municipal.json.gz", {
        "url": fontes.SINISA_GESTAO_MUNICIPAL,
        "sha256": sha256(b),
        "capturado_em": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "abas": dados,
    })


def sidra(caminho):
    url = fontes.SIDRA + caminho
    b = baixar(url, timeout=180)
    return url, sha256(b), json.loads(b)


def coletar_censo():
    capturado = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    saida = {"capturado_em": capturado, "tabelas": {}}
    for chave, t in fontes.CENSO.items():
        meta_url = fontes.SIDRA_METADADOS.format(tabela=t["tabela"])
        meta = json.loads(baixar(meta_url))
        cats = {}
        for c in meta["classificacoes"]:
            if c["id"] == t["classificacao"]:
                cats = {str(k["id"]): {"nome": k["nome"], "nivel": k.get("nivel")} for k in c["categorias"]}
        valores = {}
        consultas = []
        niveis = [("n1", "all"), ("n2", "all"), ("n3", "all")] + [("n6", f"in n3 {uf}") for uf in UFS]
        for nivel, filtro in niveis:
            caminho = (f"/t/{t['tabela']}/{nivel}/{urllib.parse.quote(filtro)}/v/{t['variavel']}/p/{t['ano']}"
                       f"/c{t['classificacao']}/all{t['fixas']}/h/n/f/c")
            url, h, linhas = sidra(caminho)
            consultas.append({"url": url, "sha256": h, "linhas": len(linhas)})
            for l in linhas:
                v = l["V"]
                valor = int(v) if re.fullmatch(r"-?\d+", v) else v  # "-" (zero absoluto) e "X" (sigilo) preservados
                valores.setdefault(f"{l['NC']}:{l['D1C']}", {})[l["D4C"]] = valor
        saida["tabelas"][chave] = {**t, "titulo": meta["nome"], "categorias": cats, "consultas": consultas, "valores": valores}
        print(f"  {chave}: {len(valores)} territórios")
    # 2010, para a evolução entre censos (mesmo conceito de "rede geral como forma principal")
    for chave, tabela, classif, fixas, variavel in [
        ("agua_2010", 3217, 61, "/c1/0/c67/0/c309/0", 137),
        ("esgoto_2010", 3216, 299, "/c1/0/c125/0/c63/0/c300/0", 137),
    ]:
        meta = json.loads(baixar(fontes.SIDRA_METADADOS.format(tabela=tabela)))
        cats = {}
        for c in meta["classificacoes"]:
            if c["id"] == classif:
                cats = {str(k["id"]): {"nome": k["nome"], "nivel": k.get("nivel")} for k in c["categorias"]}
        valores, consultas = {}, []
        niveis = [("n1", "all"), ("n2", "all"), ("n3", "all")] + [("n6", f"in n3 {uf}") for uf in UFS]
        for nivel, filtro in niveis:
            caminho = f"/t/{tabela}/{nivel}/{urllib.parse.quote(filtro)}/v/{variavel}/p/2010/c{classif}/all{fixas}/h/n/f/c"
            url, h, linhas = sidra(caminho)
            consultas.append({"url": url, "sha256": h, "linhas": len(linhas)})
            dim = None
            for l in linhas:
                if dim is None:
                    dim = next(k for k in l if k.startswith("D") and k.endswith("C") and str(l[k]) in cats)
                v = l["V"]
                valores.setdefault(f"{l['NC']}:{l['D1C']}", {})[l[dim]] = int(v) if re.fullmatch(r"-?\d+", v) else v
        saida["tabelas"][chave] = {"tabela": tabela, "variavel": variavel, "classificacao": classif, "fixas": fixas,
                                   "ano": 2010, "titulo": meta["nome"], "categorias": cats,
                                   "consultas": consultas, "valores": valores}
        print(f"  {chave}: {len(valores)} territórios")
    # população estimada 2024 (mesma referência do DFE0001 do SINISA)
    p = fontes.POPULACAO_ESTIMADA
    valores, consultas = {}, []
    for nivel in ("n1/all", "n2/all", "n3/all", "n6/all"):
        url, h, linhas = sidra(f"/t/{p['tabela']}/{nivel}/v/{p['variavel']}/p/{p['ano']}/h/n/f/c")
        consultas.append({"url": url, "sha256": h, "linhas": len(linhas)})
        for l in linhas:
            valores[f"{l['NC']}:{l['D1C']}"] = int(l["V"]) if re.fullmatch(r"\d+", l["V"]) else l["V"]
    saida["populacao_2024"] = {**p, "consultas": consultas, "valores": valores}
    # nomes oficiais dos municípios
    b = baixar(fontes.IBGE_MUNICIPIOS)
    mun = json.loads(b)
    saida["municipios"] = {
        "url": fontes.IBGE_MUNICIPIOS, "sha256": sha256(b),
        "lista": sorted(
            [{"id": str(m["id"]), "nome": m["nome"],
              "uf": (m.get("microrregiao") or {}).get("mesorregiao", {}).get("UF", {}).get("sigla")
              or (m.get("regiao-imediata") or {}).get("regiao-intermediaria", {}).get("UF", {}).get("sigla")}
             for m in mun], key=lambda x: x["id"]),
    }
    gravar("censo.json.gz", saida)


def coletar_cadunico():
    """Famílias no CadÚnico por faixa de renda, por município, nos meses de fontes.CADUNICO."""
    capturado = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    c = fontes.CADUNICO
    saida = {"capturado_em": capturado, "fonte": fontes.MISOCIAL, "campos": c["campos"], "meses": {}}
    for mes in c["meses"]:
        url = (f"{fontes.MISOCIAL}?q=anomes_s:{mes}&fl={','.join(c['campos'])}"
               "&rows=6000&wt=json&sort=codigo_ibge%20asc")
        b = baixar(url, timeout=180)
        docs = json.loads(b)["response"]["docs"]
        saida["meses"][mes] = {
            "url": url, "sha256": sha256(b), "registros": len(docs),
            "municipios": {d["codigo_ibge"]: {k: v for k, v in d.items() if k not in ("codigo_ibge", "anomes_s")} for d in docs},
        }
        print(f"  CadÚnico {mes}: {len(docs)} municípios")
    gravar("cadunico.json.gz", saida)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--so-sidra", action="store_true")
    ap.add_argument("--so-sinisa", action="store_true")
    ap.add_argument("--zips", help="diretório com os ZIPs do SINISA já baixados")
    ap.add_argument("--so-cadunico", action="store_true")
    a = ap.parse_args()
    if a.so_cadunico:
        coletar_cadunico()
        return
    if not a.so_sidra:
        coletar_sinisa(a.zips)
        coletar_gestao_municipal(os.path.join(a.zips, fontes.SINISA_GESTAO_MUNICIPAL.rsplit("/", 1)[1]) if a.zips else None)
    if not a.so_sinisa:
        coletar_censo()


if __name__ == "__main__":
    main()
