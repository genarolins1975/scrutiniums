"""Audit the inspected MS capture/recapture release; never promote health scores.

The frozen review uses stdlib only. Re-extraction from official XLSX files
requires openpyxl. A source revision or a new municipal release needs a new
explicit contract, not a silent change of the geographic unit.
"""
import argparse
import hashlib
import json
import math
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
SNAPSHOT = Path(__file__).parent / 'seed/saude/ms-captura-recaptura-2026-10-12.json'
SNAPSHOT_SHA256 = '48056f0747caf31fdace39c63238111ee4cd191e370f13345c069a7261cca937'
BASE = 'https://svs.aids.gov.br/daent/acesso-a-informacao/acoes-e-programas/captura-recaptura/mortalidade/'
YEARS = [2021, 2022, 2023]
FILES = {
    'cobertura-SIM.xlsx': {
        'sha256': 'dd3cd353f48d63ec589e67460fead6e85cb578c43b3f1d9916acfa3f4406636c',
        'sheets': {
            'Cobertura SIM Infantil': ('infant_registration_coverage', 'percent', list(range(2015, 2024))),
            'Cobertura SIM Total': ('all_age_registration_coverage', 'percent', list(range(2015, 2024))),
        },
    },
    'taxa-mortalidade-infantil.xlsx': {
        'sha256': 'a8cfccad239f134747e1fbfd41051f079dc267a6cb0d736812e69a349a9f0706',
        'sheets': {'TMI': ('infant_mortality_rate', 'per_1000_live_births', list(range(2000, 2025)))},
    },
}
STATES = {
    'Rondônia', 'Acre', 'Amazonas', 'Roraima', 'Pará', 'Amapá', 'Tocantins',
    'Maranhão', 'Piauí', 'Ceará', 'Rio Grande do Norte', 'Paraíba', 'Pernambuco',
    'Alagoas', 'Sergipe', 'Bahia', 'Minas Gerais', 'Espírito Santo',
    'Rio de Janeiro', 'São Paulo', 'Paraná', 'Santa Catarina', 'Rio Grande do Sul',
    'Mato Grosso do Sul', 'Mato Grosso', 'Goiás', 'Distrito Federal',
}
REGIONS = {'Norte', 'Nordeste', 'Sudeste', 'Sul', 'Centro-Oeste'}
TERRITORIES = STATES | REGIONS | {'Brasil'}


def geography(name):
    if name in STATES:
        return 'state'
    if name in REGIONS:
        return 'region'
    if name == 'Brasil':
        return 'country'
    raise ValueError('Unexpected territory; review the source contract: ' + str(name))


def audit(snapshot):
    """Validate this frozen state-level release. No name-based municipal join."""
    if snapshot.get('version') != 'ms-capture-recapture-scope-v1':
        raise ValueError('Unknown snapshot version')
    if snapshot.get('selected_years') != YEARS:
        raise ValueError('Health window mismatch')
    sources = snapshot.get('sources', [])
    if len(sources) != len(FILES) or {s['file'] for s in sources} != set(FILES):
        raise ValueError('Missing or duplicated workbook')
    summaries = []
    for source in sources:
        spec = FILES[source['file']]
        if source['sha256'] != spec['sha256'] or source['url'] != BASE + source['file']:
            raise ValueError('Source identity changed; review before replacing evidence')
        tables = source['tables']
        if len(tables) != len(spec['sheets']) or {t['sheet'] for t in tables} != set(spec['sheets']):
            raise ValueError('Worksheet inventory changed')
        for table in tables:
            metric, unit, years = spec['sheets'][table['sheet']]
            if (table['geography_header'] != 'UF' or table['available_years'] != years
                    or table['metric'] != metric or table['unit'] != unit):
                raise ValueError('Unexpected metric, unit, years or geographic header')
            rows = table['rows']
            if len(rows) != 33 or {r['name'] for r in rows} != TERRITORIES:
                raise ValueError('Territory inventory changed or duplicated')
            if len({r['source_row'] for r in rows}) != len(rows):
                raise ValueError('Duplicated source row')
            for row in rows:
                if row['geography_level'] != geography(row['name']):
                    raise ValueError('Geographic level mismatch; never substitute a state for a city')
                if set(row['values']) != {str(y) for y in YEARS}:
                    raise ValueError('Missing or unexpected year')
                for value in row['values'].values():
                    if (type(value) not in (int, float) or not math.isfinite(value)
                            or value < 0 or (unit == 'percent' and value > 100)):
                        raise ValueError('Missing, nonfinite or out-of-domain value')
            summaries.append({'file': source['file'], 'sheet': table['sheet'],
                              'metric': metric, 'territories': len(rows),
                              'levels': ['country', 'region', 'state'],
                              'selected_years': YEARS, 'selected_cells': len(rows) * len(YEARS)})
    return {
        'version': 'irc-health-capture-recapture-review-v1',
        'review_date': '2026-10-12',
        'required_geography': 'municipality', 'required_years': YEARS,
        'tables': summaries,
        'expected_municipal_infant_coverage_cells': 26 * len(YEARS),
        'available_municipal_infant_coverage_cells_in_inspected_files': 0,
        'interpretation': 'Zero eligible evidence cells does not mean zero coverage of the registers.',
        'decision': 'state_context_only', 'health_score': None,
        'score_eligible': False, 'correction_applied': False, 'publication_ready': False,
        'next_requirement': 'Municipal infant-specific coverage or corrected events/rates with documented denominators and a compatible temporal contract.',
    }


def extract(downloads):
    """Read all workbook sheets/labels, freeze 2021-2023 values and source cells."""
    import openpyxl
    from openpyxl.utils import get_column_letter
    output = {'version': 'ms-capture-recapture-scope-v1', 'retrieved_at': '2026-10-12',
              'selected_years': YEARS, 'sources': []}
    for filename, spec in FILES.items():
        path = Path(downloads) / filename
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        if digest != spec['sha256']:
            raise ValueError('Downloaded release changed: ' + filename)
        book = openpyxl.load_workbook(path, data_only=True)
        if set(book.sheetnames) != set(spec['sheets']):
            raise ValueError('Workbook sheets changed')
        source = {'file': filename, 'url': BASE + filename, 'sha256': digest, 'tables': []}
        for sheet in book:
            metric, unit, expected_years = spec['sheets'][sheet.title]
            matrix = list(sheet.values)
            header = matrix[0]
            years = [v for v in header[1:] if type(v) is int]
            if header[0] != 'UF' or years != expected_years:
                raise ValueError('Workbook header changed')
            columns = {y: header.index(y) for y in YEARS}
            table = {'sheet': sheet.title, 'sheet_state': sheet.sheet_state,
                     'geography_header': header[0], 'metric': metric, 'unit': unit,
                     'available_years': years,
                     'source_columns': {str(y): get_column_letter(columns[y] + 1) for y in YEARS},
                     'rows': [], 'notes': []}
            for number, row in enumerate(matrix[1:], 2):
                # Numeric data may not be silently discarded as a footnote.
                has_data = any(v is not None for v in row[1:len(years) + 1])
                if row[0] in TERRITORIES or has_data:
                    level = geography(row[0])
                    table['rows'].append({'source_row': number, 'name': row[0],
                                          'geography_level': level,
                                          'values': {str(y): row[columns[y]] for y in YEARS}})
                elif row[0] is not None:
                    table['notes'].append({'cell': 'A' + str(number), 'text': str(row[0])})
            source['tables'].append(table)
        output['sources'].append(source)
    audit(output)
    return output


def write_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def review_frozen():
    raw = SNAPSHOT.read_bytes()
    if hashlib.sha256(raw).hexdigest() != SNAPSHOT_SHA256:
        raise ValueError('Frozen evidence changed; review the source before updating its hash')
    result = audit(json.loads(raw))
    result['input_sha256'] = SNAPSHOT_SHA256
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--downloads', type=Path, help='Re-extract XLSX into --output instead of reviewing the frozen snapshot')
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    if args.downloads:
        result = extract(args.downloads)
    else:
        result = review_frozen()
    write_json(args.output, result)


if __name__ == '__main__':
    main()
