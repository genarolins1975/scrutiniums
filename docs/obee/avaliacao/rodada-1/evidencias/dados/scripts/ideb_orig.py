import openpyxl, json
from lib import *
wb=openpyxl.load_workbook('inep/x/divulgacao_anos_iniciais_municipios_2025/divulgacao_anos_iniciais_municipios_2025.xlsx',read_only=True,data_only=True)
print(wb.sheetnames)
ws=wb[wb.sheetnames[0]]
rows=ws.iter_rows(values_only=True)
hdr=None
data=[]
for i,r in enumerate(rows):
    if hdr is None and r and 'CO_MUNICIPIO' in [str(x) for x in r]: hdr=[str(x) for x in r]; continue
    if hdr: data.append(dict(zip(hdr,r)))
print(len(data), hdr[:12])
cods={c['cod_ibge'] for c in CAPS}
mun=[d for d in data if d.get('CO_MUNICIPIO') is not None and int(d['CO_MUNICIPIO']) in cods and d['REDE']=='Municipal']
print(len(mun))
bad=0;n=0;nd=0
for d in mun:
    cod=int(d['CO_MUNICIPIO'])
    for a in [2005,2007,2009,2011,2013,2015,2017,2019,2021,2023,2025]:
        v=d.get(f'VL_OBSERVADO_{a}')
        o=OBS.get(('edu.ideb.rede_municipal',cod,a,'anos_iniciais','ideb'))
        if isinstance(v,(int,float)):
            n+=1
            if o['valor'] is None or abs(float(v)-o['valor'])>1e-9: bad+=1; print('DIVERGE',cod,a,v,o['valor'],o['status'])
        else:
            nd+=1
            if o['status']=='OBSERVADO': bad+=1; print('DIVERGE nd',cod,a,v,o['status'])
print('comparações com valor',n,'sem valor',nd,'divergências',bad)
