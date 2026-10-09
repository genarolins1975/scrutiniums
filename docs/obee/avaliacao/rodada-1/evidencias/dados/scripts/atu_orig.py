import openpyxl
from lib import *
def le(path):
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    ws=wb[wb.sheetnames[0]]; hdr=None; out=[]
    for r in ws.iter_rows(values_only=True):
        if hdr is None and r and 'CO_MUNICIPIO' in [str(x) for x in r]: hdr=[str(x) for x in r]; continue
        if hdr: out.append(dict(zip(hdr,r)))
    return out
cods={c['cod_ibge'] for c in CAPS}
A=[d for d in le('inep/atu/ATU_2025_MUNICIPIOS/ATU_MUNICIPIOS_2025.xlsx') if d.get('CO_MUNICIPIO') and int(d['CO_MUNICIPIO']) in cods and d['NO_CATEGORIA']=='Total' and d['NO_DEPENDENCIA']=='Municipal']
Rr=[d for d in le('inep/rend/tx_rend_municipios_2025/tx_rend_municipios_2025.xlsx') if d.get('CO_MUNICIPIO') and int(d['CO_MUNICIPIO']) in cods and d['NO_CATEGORIA']=='Total' and d['NO_DEPENDENCIA']=='Municipal']
print(len(A),len(Rr))
bad=n=0
for d in A:
    cod=int(d['CO_MUNICIPIO'])
    for e,col in {'creche':'CRE_CAT_0','pre_escola':'PRE_CAT_0','anos_iniciais':'FUN_AI_CAT_0','anos_finais':'FUN_AF_CAT_0'}.items():
        v=d[col]; o=OBS[('edu.atu.rede_municipal',cod,2025,e,None)]
        if isinstance(v,(int,float)):
            n+=1
            if o['valor'] is None or abs(float(v)-o['valor'])>1e-9: bad+=1; print('DIV atu',cod,e,v,o['valor'])
        elif o['status']=='OBSERVADO': bad+=1; print('DIV atu nd',cod,e,v)
for d in Rr:
    cod=int(d['CO_MUNICIPIO'])
    for e,col in {'anos_iniciais':'1_CAT_FUN_AI','anos_finais':'1_CAT_FUN_AF'}.items():
        v=d[col]; o=OBS[('edu.aprovacao.rede_municipal',cod,2025,e,None)]
        if isinstance(v,(int,float)):
            n+=1
            if o['valor'] is None or abs(float(v)-o['valor'])>1e-9: bad+=1; print('DIV aprov',cod,e,v,o['valor'])
        elif o['status']=='OBSERVADO': bad+=1; print('DIV aprov nd',cod,e,v)
print('comparações',n,'divergências',bad)
