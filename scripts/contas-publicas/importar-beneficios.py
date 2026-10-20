"""Gastos tributários da União (Receita Federal, DGT Bases Efetivas): integração por edição, nunca no build.

Uso:
  python3 scripts/contas-publicas/importar-beneficios.py              # baixa a planilha, confere o hash e refaz tudo
  python3 scripts/contas-publicas/importar-beneficios.py --xlsx X     # mesma coisa a partir de um arquivo local
  python3 scripts/contas-publicas/importar-beneficios.py --verificar  # refaz o snapshot dos quadros guardados, sem rede

Fonte: "DGT Bases Efetivas 2023 · Série 2021 a 2026 · Quadros" (Receita Federal, publicada em 07/04/2026 e
atualizada em 21/08/2026). Estimativas oficiais do que a União deixou de arrecadar com cada benefício
tributário: 2021 a 2023 calculados com dados efetivos (declarações e escriturações); 2024 a 2026 são projeções
da mesma edição. Valores em R$ correntes.

Duas etapas, para que o teste refaça o snapshot sem rede e sem openpyxl:
  1. extrair: lê da planilha só os quadros usados (V, VI, X e XXVI) e grava as células em
     data/contas-publicas/beneficios-dgt-2023/quadros.json, com o hash da planilha no manifesto;
  2. transformar: monta o snapshot e confere as somas publicadas (tributos, funções e modalidades contra o total
     do quadro, e o total de 2023 do quadro XXVI contra o do quadro VI). Qualquer divergência reprova e nada é
     gravado.

Edição nova (a Receita publica uma por ano) muda o hash: o importador recusa e a integração é revisada à mão
(OPERACOES.md), com novo diretório de quadros.
"""
import argparse
import datetime as dt
import hashlib
import io
import json
import math
import pathlib
import urllib.request

ROOT = pathlib.Path(__file__).resolve().parents[2]
EDICAO = 'DGT Bases Efetivas 2023 · Série 2021 a 2026'
URL = ('https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/publicacoes/relatorios/renuncia/'
       'gastos-tributarios-bases-efetivas/dgt-bases-efetivas-2023-serie-2021-a-2026-quadros.xlsx/@@download/file')
PAGINA = ('https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/publicacoes/relatorios/renuncia/'
          'gastos-tributarios-bases-efetivas')
SHA256 = '38b28ceefc7274b0e2e68d436eb7c756fb8cfac266b0228ba0fc8a544a1e22bb'
PUBLICACAO = '2026-04-07'
ATUALIZACAO_PAGINA = '2026-08-21'
ORIGEM = ROOT / 'data/contas-publicas/beneficios-dgt-2023'
DESTINO = ROOT / 'src/lib/contas-publicas/snapshots/beneficios.json'
QUADROS = ('Q V', 'Q VI', 'Q X', 'Q XXVI')
ANO_BASE = 2023
TOLERANCIA = 1.0  # R$; a planilha soma em ponto flutuante


def check(ok, mensagem):
    if not ok:
        raise ValueError(mensagem)


def _num(v, onde):
    check(isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v), f'{onde}: valor não numérico ({v!r})')
    return float(v)


def extrair(xlsx):
    """Células dos quadros usados, sem linhas vazias; datas viram texto ISO."""
    import openpyxl  # só a extração lê a planilha
    wb = openpyxl.load_workbook(io.BytesIO(xlsx), data_only=True, read_only=True)
    quadros = {}
    for nome in QUADROS:
        check(nome in wb.sheetnames, f'Quadro ausente: {nome}')
        linhas = []
        for r in wb[nome].iter_rows(values_only=True):
            celulas = [c.isoformat() if isinstance(c, (dt.date, dt.datetime)) else c for c in r[:8]]
            while celulas and celulas[-1] is None:
                celulas.pop()
            if celulas:
                linhas.append(celulas)
        quadros[nome] = linhas
    return quadros


def _ate_total(linhas, inicio, onde):
    corpo = []
    for l in linhas[inicio:]:
        if l[0] == 'TOTAL':
            return corpo, l
        corpo.append(l)
    raise ValueError(f'{onde}: linha TOTAL ausente')


def transformar(quadros, manifesto):
    check(manifesto.get('sha256') == SHA256, 'Planilha diferente da edição integrada: revisar a edição nova à mão')
    # Quadro VI: por tributo em 2023, com as razões publicadas ao PIB e à arrecadação
    q6 = quadros['Q VI']
    check(any('BASES EFETIVAS 2023' in str(l[0]) for l in q6[:3]), 'Quadro VI não é da base efetiva de 2023')
    i = next(k for k, l in enumerate(q6) if l[0] == 'TRIBUTO')
    corpo, total6 = _ate_total(q6, i + 2, 'Quadro VI')
    tributos = [{'tributo': l[0], 'valor': _num(l[1], l[0]), 'pctPIB': _num(l[2], l[0]), 'pctArrecadacao': _num(l[3], l[0])} for l in corpo]
    total = _num(total6[1], 'Quadro VI total')
    arrec = next(l for l in q6 if l[0] == 'ARRECADAÇÃO*')
    pib = next(l for l in q6 if l[0] == 'PIB')
    arrecadacao, pib_valor = _num(arrec[1], 'arrecadação'), _num(pib[1], 'PIB')
    check(abs(sum(t['valor'] for t in tributos) - total) <= TOLERANCIA, 'Quadro VI: tributos não somam o total')
    check(abs(100 * total / pib_valor - _num(total6[2], 'GT/PIB')) <= 1e-9, 'Quadro VI: razão ao PIB não confere')
    check(abs(100 * total / arrecadacao - _num(total6[3], 'GT/arrecadação')) <= 1e-9, 'Quadro VI: razão à arrecadação não confere')
    # Quadro V: por função orçamentária (a classificação da própria Receita, aproximada da funcional do orçamento)
    q5 = quadros['Q V']
    i = next(k for k, l in enumerate(q5) if l[0] == 'FUNÇÃO ORÇAMENTÁRIA')
    corpo, total5 = _ate_total(q5, i + 1, 'Quadro V')
    funcoes = [{'funcao': l[0], 'valor': _num(l[1], l[0])} for l in corpo]
    check(abs(_num(total5[1], 'Quadro V total') - total) <= TOLERANCIA, 'Quadro V: total difere do Quadro VI')
    check(abs(sum(f['valor'] for f in funcoes) - total) <= TOLERANCIA, 'Quadro V: funções não somam o total')
    # Quadro X: por modalidade de benefício
    q10 = quadros['Q X']
    i = next(k for k, l in enumerate(q10) if l[0] == 'GASTO TRIBUTÁRIO')
    corpo, total10 = _ate_total(q10, i + 1, 'Quadro X')
    modalidades = [{'gasto': str(l[0]).strip(), 'valor': _num(l[1], l[0])} for l in corpo]
    check(abs(_num(total10[1], 'Quadro X total') - total) <= TOLERANCIA, 'Quadro X: total difere do Quadro VI')
    check(abs(sum(m['valor'] for m in modalidades) - total) <= TOLERANCIA, 'Quadro X: modalidades não somam o total')
    check(len({m['gasto'] for m in modalidades}) == len(modalidades), 'Quadro X: modalidade repetida')
    # Quadro XXVI: série 2021 a 2026, efetivo e projeção
    q26 = quadros['Q XXVI']
    i = next(k for k, l in enumerate(q26) if l[0] == 'TRIBUTO / GASTO TRIBUTÁRIO')
    cab, anos_l = q26[i], q26[i + 1]
    check(cab[1] == 'EFETIVO' and cab[4] == 'PROJEÇÃO', 'Quadro XXVI: cabeçalho efetivo/projeção mudou')
    anos = [int(a) for a in anos_l[1:7]]
    check(anos == list(range(anos[0], anos[0] + 6)) and anos[2] == ANO_BASE, 'Quadro XXVI: anos inesperados')
    linha = {l[0]: l for l in q26[i + 2:] if isinstance(l[0], str)}
    for chave in ('TOTAL', 'Receita Administrada - RFB', 'PIB', 'GT / PIB', 'GT / Arrecadação'):
        check(chave in linha, f'Quadro XXVI: linha ausente ({chave})')
    serie = []
    for k, ano in enumerate(anos):
        valor = _num(linha['TOTAL'][k + 1], f'total {ano}')
        pib_ano = _num(linha['PIB'][k + 1], f'PIB {ano}')
        arr_ano = _num(linha['Receita Administrada - RFB'][k + 1], f'arrecadação {ano}')
        check(abs(valor / pib_ano - _num(linha['GT / PIB'][k + 1], 'GT/PIB')) <= 1e-12, f'Quadro XXVI: GT/PIB de {ano} não confere')
        check(abs(valor / arr_ano - _num(linha['GT / Arrecadação'][k + 1], 'GT/arrecadação')) <= 1e-12, f'Quadro XXVI: GT/arrecadação de {ano} não confere')
        por_tributo = sum(_num(linha[t['tributo']][k + 1], t['tributo']) for t in tributos if t['tributo'] in linha)
        check(abs(por_tributo - valor) <= TOLERANCIA, f'Quadro XXVI: tributos não somam o total de {ano}')
        serie.append({'ano': ano, 'natureza': 'efetivo' if k < 3 else 'projecao', 'valor': valor,
                      'pctPIB': 100 * valor / pib_ano, 'pctArrecadacao': 100 * valor / arr_ano,
                      'pib': pib_ano, 'arrecadacao': arr_ano})
    check(all(t['tributo'] in linha for t in tributos), 'Quadro XXVI: tributo do Quadro VI ausente da série')
    base = next(s for s in serie if s['ano'] == ANO_BASE)
    check(abs(base['valor'] - total) <= TOLERANCIA, 'Quadro XXVI: total de 2023 difere do Quadro VI')
    return {
        'versao': 1,
        'transformacao': 'importar-beneficios.py v1',
        'edicao': EDICAO,
        'fonte': {'nome': 'Receita Federal · Demonstrativo dos Gastos Tributários, Bases Efetivas', 'url': URL, 'pagina': PAGINA,
                  'sha256': SHA256, 'publicacao': PUBLICACAO, 'atualizacaoPagina': ATUALIZACAO_PAGINA,
                  'quadros': 'V (por função), VI (por tributo), X (por modalidade) e XXVI (série 2021 a 2026)'},
        'captura': manifesto['captura'],
        'unidade': 'R$ correntes',
        'anoBase': ANO_BASE,
        'perimetro': 'Tributos federais administrados pela Receita Federal; arrecadação sem a contribuição do servidor (CPSS)',
        'natureza': 'Estimativa oficial de receita que deixou de ser arrecadada; 2021 a 2023 com dados efetivos, 2024 a 2026 projeções da mesma edição',
        'total': total,
        'pib': pib_valor,
        'arrecadacao': arrecadacao,
        'serie': serie,
        'tributos': sorted(tributos, key=lambda t: -t['valor']),
        'funcoes': sorted(funcoes, key=lambda f: -f['valor']),
        'modalidades': sorted(modalidades, key=lambda m: -m['valor']),
        'conferencia': 'Tributos, funções e modalidades somam o total do Quadro VI com tolerância de R$ 1; o total de 2023 do Quadro XXVI é o mesmo; as razões ao PIB e à arrecadação publicadas foram recalculadas.',
    }


def main():
    ap = argparse.ArgumentParser()
    g = ap.add_mutually_exclusive_group()
    g.add_argument('--xlsx', type=pathlib.Path, help='planilha já baixada')
    g.add_argument('--verificar', action='store_true', help='refaz o snapshot dos quadros guardados, sem rede')
    a = ap.parse_args()
    if a.verificar:
        quadros = json.loads((ORIGEM / 'quadros.json').read_text())
        manifesto = json.loads((ORIGEM / 'manifesto.json').read_text())
    else:
        xlsx = a.xlsx.read_bytes() if a.xlsx else urllib.request.urlopen(URL, timeout=120).read()
        manifesto = {'sha256': hashlib.sha256(xlsx).hexdigest(), 'url': URL, 'captura': dt.datetime.now(dt.timezone.utc).isoformat(), 'bytes': len(xlsx)}
        quadros = extrair(xlsx)
    snapshot = transformar(quadros, manifesto)  # valida antes de gravar
    if not a.verificar:
        ORIGEM.mkdir(parents=True, exist_ok=True)
        (ORIGEM / 'quadros.json').write_text(json.dumps(quadros, ensure_ascii=False, separators=(',', ':')) + '\n')
        (ORIGEM / 'manifesto.json').write_text(json.dumps(manifesto, ensure_ascii=False, indent=2) + '\n')
    DESTINO.write_text(json.dumps(snapshot, ensure_ascii=False, indent=1) + '\n')
    print(f"{EDICAO}: R$ {snapshot['total'] / 1e9:.1f} bi em {ANO_BASE}; {len(snapshot['modalidades'])} modalidades; série {snapshot['serie'][0]['ano']} a {snapshot['serie'][-1]['ano']}")


if __name__ == '__main__':
    main()
